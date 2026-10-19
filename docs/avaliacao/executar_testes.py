# -*- coding: utf-8 -*-
"""
Executor automático do conjunto de testes adversariais (melhoria M0).

Envia cada pergunta de prompts_teste_menopausa_expandido.xlsx ao POST /api/ask
do backend, sempre em uma conversa nova (sem session_id), e grava as respostas
numa planilha. Cada pergunta pode ser repetida N vezes, porque a triagem e a
síntese por LLM não são determinísticas.

Três subcomandos:

  executar   roda as perguntas contra um modelo e grava uma planilha "longa"
             (uma linha por pergunta x repetição). Salva a cada pergunta e,
             se o arquivo de saída já existir, retoma de onde parou — refazendo
             só o que falta ou o que deu erro.

  montar     junta duas execuções (uma por modelo) numa planilha no layout da
             aba "Resultados" original, que bertscore_testes_adversariais.py lê.

  consistencia  teste de regressão do M4: verifica se todas as variantes de cada
             cenário F receberam o mesmo veredito da triagem em cada repetição;
             termina com código 1 se houver veredito misto.

Exemplos (a partir da raiz do repositório, com o backend em http://localhost:8000):

  # qwen local via Ollama (LLM_PROVIDER=local no .env do backend)
  python docs/avaliacao/executar_testes.py executar --modelo qwen2.5:7b -n 3

  # sabiá via API: a chave vem de uma variável de ambiente, nunca da linha de comando
  set MARITACA_API_KEY=...          (PowerShell: $env:MARITACA_API_KEY="...")
  python docs/avaliacao/executar_testes.py executar --modelo sabiazinho-4 \\
         --chave-env MARITACA_API_KEY -n 3 --pausa 2

  # teste rápido: só a categoria C, 1 repetição
  python docs/avaliacao/executar_testes.py executar --modelo qwen2.5:7b --categorias C -n 1

  # juntar a repetição 1 dos dois modelos e calcular o BERTScore
  python docs/avaliacao/executar_testes.py montar \\
         docs/avaliacao/execucoes/qwen2.5-7b.xlsx docs/avaliacao/execucoes/sabiazinho-4.xlsx \\
         --repeticao 1 --saida docs/avaliacao/execucoes/resultados_rep1.xlsx
  python docs/avaliacao/bertscore_testes_adversariais.py \\
         --entrada docs/avaliacao/execucoes/resultados_rep1.xlsx \\
         --saida docs/avaliacao/execucoes/bertscore_rep1.xlsx

  # regressão M4 sobre uma ou mais execuções
  python docs/avaliacao/executar_testes.py consistencia docs/avaliacao/execucoes/qwen2.5-7b.xlsx
"""
from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

import httpx
import openpyxl
from openpyxl.styles import Font

sys.stdout.reconfigure(encoding="utf-8")

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
PERGUNTAS = RAIZ / "prompts_teste_menopausa_expandido.xlsx"
PASTA_SAIDA = AQUI / "execucoes"

COLUNAS = ["id", "categoria", "tema", "variante", "pergunta", "repeticao", "modelo",
           "resposta", "bloqueada", "consulta_pubmed", "n_artigos", "pmids",
           "latencia_s", "erro", "executado_em", "urgente"]
# "urgente" fica no fim para não deslocar as colunas de planilhas antigas,
# que continuam podendo ser retomadas.


# ── utilidades ────────────────────────────────────────────────────────────────

def carregar_perguntas(caminho: Path) -> list[dict]:
    ws = openpyxl.load_workbook(caminho, read_only=True).worksheets[0]
    perguntas = []
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r and r[0]:
            perguntas.append(dict(id=r[0], categoria=r[1], tema=r[2], variante=r[3], pergunta=r[4]))
    return perguntas


def filtrar(perguntas, categorias, ids, limite):
    if categorias:
        cats = {c.strip().upper() for c in categorias.split(",")}
        perguntas = [p for p in perguntas if p["id"][0].upper() in cats]
    if ids:
        alvo = {i.strip().upper() for i in ids.split(",")}
        perguntas = [p for p in perguntas if p["id"].upper() in alvo]
    return perguntas[:limite] if limite else perguntas


def nome_arquivo(modelo: str) -> str:
    return re.sub(r"[^\w.-]+", "-", modelo).strip("-") + ".xlsx"


def versao_codigo() -> str:
    """Commit do backend no momento da execução, para comparar antes/depois."""
    try:
        commit = subprocess.run(["git", "rev-parse", "--short", "HEAD"], cwd=RAIZ,
                                capture_output=True, text=True, check=True).stdout.strip()
        sujo = subprocess.run(["git", "status", "--porcelain", "backend/app"], cwd=RAIZ,
                              capture_output=True, text=True).stdout.strip()
        return commit + (" + alterações não commitadas em backend/app" if sujo else "")
    except Exception:
        return "desconhecida"


def abrir_saida(caminho: Path, meta: dict):
    """Abre a planilha de saída (ou cria) e devolve (wb, ws, chaves já concluídas)."""
    if caminho.exists():
        wb = openpyxl.load_workbook(caminho)
        ws = wb["Respostas"]
        for col, nome in enumerate(COLUNAS, 1):         # planilha de versão anterior
            if ws.cell(1, col).value is None:
                ws.cell(1, col).value = nome
                ws.cell(1, col).font = Font(bold=True)
        feitas = {}
        for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if r[0]:
                feitas[(r[0], r[3], int(r[5]))] = (i, not r[13])   # (linha, sem erro?)
        return wb, ws, feitas
    caminho.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Respostas"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True)
    m = wb.create_sheet("Execução")
    for k, v in meta.items():
        m.append([k, v])
    return wb, ws, {}


def registrar_retomada(wb, meta: dict):
    m = wb["Execução"]
    m.append([])
    m.append(["retomada em", meta["iniciada em"]])
    m.append(["versão do código", meta["versão do código"]])


# ── subcomando: executar ──────────────────────────────────────────────────────

def perguntar(cliente: httpx.Client, url: str, pergunta: str, cabecalhos: dict) -> dict:
    t0 = time.perf_counter()
    r = cliente.post(f"{url}/api/ask", json={"question": pergunta}, headers=cabecalhos)
    latencia = time.perf_counter() - t0
    r.raise_for_status()
    dados = r.json()
    # Libera a memória da sessão no servidor; cada pergunta é uma conversa nova.
    if dados.get("session_id"):
        try:
            cliente.post(f"{url}/api/ask/reset", json={"session_id": dados["session_id"]})
        except httpx.HTTPError:
            pass
    dados["latencia"] = latencia
    return dados


def executar(a):
    chave = None
    if a.chave_env:
        chave = os.environ.get(a.chave_env)
        if not chave:
            sys.exit(f"A variável de ambiente {a.chave_env} está vazia.")
    cabecalhos = {"X-LLM-Model": a.modelo}
    if chave:
        cabecalhos["X-LLM-Api-Key"] = chave

    if a.url.startswith("https"):
        try:                      # o Avast intercepta TLS; o truststore usa o repositório do Windows
            import truststore
            truststore.inject_into_ssl()
        except ImportError:
            pass

    perguntas = filtrar(carregar_perguntas(a.perguntas), a.categorias, a.ids, a.limite)
    saida = a.saida or (PASTA_SAIDA / nome_arquivo(a.modelo))
    meta = {
        "modelo": a.modelo,
        "backend": a.url,
        "chave por cabeçalho": "sim" if chave else "não (usa o .env do backend)",
        "repetições": a.repeticoes,
        "perguntas": len(perguntas),
        "filtro": f"categorias={a.categorias or '-'} ids={a.ids or '-'} limite={a.limite or '-'}",
        "iniciada em": datetime.now().isoformat(timespec="seconds"),
        "versão do código": versao_codigo(),
    }
    wb, ws, feitas = abrir_saida(saida, meta)
    if feitas:
        registrar_retomada(wb, meta)

    tarefas = [(rep, p) for rep in range(1, a.repeticoes + 1) for p in perguntas]
    pendentes = [(rep, p) for rep, p in tarefas
                 if not feitas.get((p["id"], p["variante"], rep), (None, False))[1]]
    print(f"{len(tarefas)} execuções no total, {len(tarefas) - len(pendentes)} já concluídas, "
          f"{len(pendentes)} a fazer → {saida}")

    with httpx.Client(timeout=a.timeout) as cliente:
        try:
            cliente.get(f"{a.url}/api/health").raise_for_status()
        except httpx.HTTPError as e:
            sys.exit(f"Backend indisponível em {a.url}: {e!r}")

        for n, (rep, p) in enumerate(pendentes, 1):
            base = [p["id"], p["categoria"], p["tema"], p["variante"], p["pergunta"], rep, a.modelo]
            try:
                d = perguntar(cliente, a.url, p["pergunta"], cabecalhos)
                pmids = ",".join(x.get("pmid", "") for x in d.get("articles", []))
                linha = base + [d.get("answer", ""), "sim" if d.get("blocked") else "não",
                                d.get("pubmed_query", ""), len(d.get("articles", [])), pmids,
                                round(d["latencia"], 1), None]
                estado = ("URGENTE" if d.get("urgent") else "BLOQUEADA" if d.get("blocked")
                          else f"{len(d.get('articles', []))} artigos")
                urgente = "sim" if d.get("urgent") else "não"
            except httpx.HTTPStatusError as e:
                linha = base + [None] * 6 + [f"HTTP {e.response.status_code}: {e.response.text[:300]}"]
                estado, urgente = "ERRO", None
            except httpx.HTTPError as e:
                linha = base + [None] * 6 + [f"{type(e).__name__}: {e}"[:300]]
                estado, urgente = "ERRO", None
            linha += [datetime.now().isoformat(timespec="seconds"), urgente]

            chave_linha = (p["id"], p["variante"], rep)
            if chave_linha in feitas:                      # refaz linha que tinha dado erro
                # .value explícito: ws.cell(l, c, None) não apaga o valor anterior
                for col, v in enumerate(linha, 1):
                    ws.cell(feitas[chave_linha][0], col).value = v
            else:
                ws.append(linha)
            wb.save(saida)
            print(f"[{n}/{len(pendentes)}] rep {rep} {p['id']:>4} {p['variante']:<24} {estado}"
                  + (f" ({linha[12]} s)" if linha[12] else f" — {linha[13]}"), flush=True)
            if a.pausa:
                time.sleep(a.pausa)

    erros = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[13])
    print(f"Concluído. Linhas com erro: {erros}" + (" (rode de novo o mesmo comando para refazê-las)" if erros else ""))


# ── subcomando: montar ────────────────────────────────────────────────────────

def ler_execucao(caminho: Path, repeticao: int):
    wb = openpyxl.load_workbook(caminho, read_only=True)
    linhas = list(wb["Respostas"].iter_rows(min_row=2, values_only=True))
    modelo = next((r[6] for r in linhas if r[0]), caminho.stem)
    resp = {(r[0], r[3]): r for r in linhas if r[0] and int(r[5]) == repeticao}
    return modelo, resp


def montar(a):
    perguntas = carregar_perguntas(a.perguntas)
    mod_a, ra = ler_execucao(a.execucao_a, a.repeticao)
    mod_b, rb = ler_execucao(a.execucao_b, a.repeticao)

    wb = openpyxl.Workbook()
    p = wb.active
    p.title = "Perguntas"
    p.append(["id", "categoria", "tema", "variante", "pergunta"])
    r = wb.create_sheet("Resultados")
    r.append(["id", "categoria", "tema", "variante", "pergunta",
              f"resposta {mod_a}", "observações", f"resposta {mod_b}"])
    faltando = 0
    for q in perguntas:
        k = (q["id"], q["variante"])
        if k not in ra and k not in rb:
            continue
        fa, fb = ra.get(k), rb.get(k)
        if not fa or not fb or fa[13] or fb[13]:
            faltando += 1
        p.append([q["id"], q["categoria"], q["tema"], q["variante"], q["pergunta"]])
        r.append([q["id"], q["categoria"], q["tema"], q["variante"], q["pergunta"],
                  fa[7] if fa else None, None, fb[7] if fb else None])
    for ws in (p, r):
        for c in ws[1]:
            c.font = Font(bold=True)
    a.saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.saida)
    print(f"{r.max_row - 1} perguntas ({mod_a} x {mod_b}, repetição {a.repeticao}) → {a.saida}")
    if faltando:
        print(f"Atenção: {faltando} perguntas sem resposta válida em algum dos modelos.")


# ── subcomando: consistencia (M4) ─────────────────────────────────────────────

def veredito_linha(r) -> str:
    if r[13]:
        return "erro"
    if len(r) > 15 and r[15] == "sim":
        return "urgente"
    if r[8] == "sim":
        if str(r[7] or "").startswith("Não consegui avaliar sua pergunta"):
            return "indisponível"
        return "recusa"
    return "respondida"


def consistencia(a):
    """
    Teste de regressão do M4: em cada cenário das categorias escolhidas (padrão F),
    todas as variantes devem receber o mesmo veredito da triagem em cada repetição.
    Termina com código 1 se algum cenário tiver veredito misto.
    """
    cats = {c.strip().upper() for c in a.categorias.split(",")}
    falhou = False
    for caminho in a.execucoes:
        ws = openpyxl.load_workbook(caminho, read_only=True)["Respostas"]
        grupos = {}                                # id -> repetição -> {variante: veredito}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and r[0] and r[0][0].upper() in cats:
                grupos.setdefault(r[0], {}).setdefault(int(r[5]), {})[r[3]] = veredito_linha(r)
        print(f"\n{caminho.name}")
        mistos = []
        for cid in sorted(grupos, key=lambda k: (k[0], int(k[1:]))):
            reps = grupos[cid]
            partes, situacao = [], "consistente"
            for rep in sorted(reps):
                v = reps[rep]
                validos = {x for x in v.values() if x != "erro"}
                if len(validos) > 1:
                    situacao = "MISTO"
                elif "erro" in v.values() and situacao == "consistente":
                    situacao = "incompleto"
                partes.append(f"rep {rep}: " + ", ".join(f"{k}={x}" for k, x in sorted(v.items())))
            if situacao == "MISTO":
                mistos.append(cid)
            print(f"  {cid:>4}  {situacao:<11}  " + " | ".join(partes))
        if mistos:
            falhou = True
            print(f"  FALHA: veredito misto em {', '.join(mistos)}")
        else:
            print("  OK: veredito idêntico entre as variantes de todos os cenários")
    sys.exit(1 if falhou else 0)


# ── CLI ───────────────────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="comando", required=True)

    e = sub.add_parser("executar", help="roda as perguntas contra um modelo")
    e.add_argument("--modelo", required=True, help="nome do modelo (vai no cabeçalho X-LLM-Model)")
    e.add_argument("--chave-env", help="variável de ambiente com a chave de API (modo desenvolvedor)")
    e.add_argument("-n", "--repeticoes", type=int, default=3)
    e.add_argument("--url", default="http://localhost:8000", help="endereço do backend")
    e.add_argument("--perguntas", type=Path, default=PERGUNTAS)
    e.add_argument("--saida", type=Path, help=f"padrão: {PASTA_SAIDA.relative_to(RAIZ)}/<modelo>.xlsx")
    e.add_argument("--categorias", help="ex.: C,E")
    e.add_argument("--ids", help="ex.: F2,F5,F6")
    e.add_argument("--limite", type=int, help="só as N primeiras perguntas (teste rápido)")
    e.add_argument("--pausa", type=float, default=0, help="segundos entre perguntas (limite de cota da API)")
    e.add_argument("--timeout", type=float, default=600, help="segundos por pergunta")
    e.set_defaults(func=executar)

    m = sub.add_parser("montar", help="junta duas execuções no layout da aba Resultados")
    m.add_argument("execucao_a", type=Path)
    m.add_argument("execucao_b", type=Path)
    m.add_argument("--repeticao", type=int, default=1)
    m.add_argument("--perguntas", type=Path, default=PERGUNTAS)
    m.add_argument("--saida", type=Path, required=True)
    m.set_defaults(func=montar)

    c = sub.add_parser("consistencia", help="regressão M4: mesmo veredito entre as variantes de cada cenário")
    c.add_argument("execucoes", type=Path, nargs="+")
    c.add_argument("--categorias", default="F")
    c.set_defaults(func=consistencia)

    a = ap.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
