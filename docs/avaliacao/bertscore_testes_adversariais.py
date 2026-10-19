# -*- coding: utf-8 -*-
"""
Similaridade semântica (BERTScore com BERTimbau) das respostas dos testes adversariais.

Duas análises, ambas só com respostas efetivamente dadas (recusas da triagem,
que começam com "Oi! 😊 Sou especializada...", e as mensagens fixas de urgência
e de triagem indisponível ficam de fora):

  1. Intramodelo: para cada id (cenário) e cada modelo, compara as respostas às
     variantes do cenário duas a duas — todas as combinações C(n, 2), cada par
     exatamente uma vez (3 variantes -> 3 pares).
  2. Entre modelos: para cada pergunta, compara a resposta do qwen2.5:7b com a
     do sabiazinho-4.

Como referência de piso, calcula também pares aleatórios de respostas de ids
diferentes (o quanto duas respostas sobre assuntos distintos já se parecem).

Uso:
    pip install torch bert-score openpyxl
    python bertscore_testes_adversariais.py                 # planilha original
    python bertscore_testes_adversariais.py --entrada X.xlsx --saida Y.xlsx

A entrada é qualquer planilha com a aba "Resultados" no layout original
(colunas F e H = respostas dos dois modelos, nomeadas "resposta <modelo>"),
como a gerada por `executar_testes.py montar`.
"""
import argparse, itertools, random, statistics, sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from bert_score import BERTScorer

sys.stdout.reconfigure(encoding="utf-8")

RAIZ = Path(__file__).resolve().parents[2]
ENTRADA = RAIZ / "prompts_teste_menopausa_expandido.xlsx"
SAIDA = Path(__file__).with_name("bertscore_testes_adversariais.xlsx")

MODELO = "neuralmind/bert-base-portuguese-cased"   # BERTimbau base
CAMADA = 9      # BERTimbau não está na tabela do bert-score; 9 é a camada usada
                # para o BERT-base multilíngue, de mesma arquitetura (12 camadas)
N_PISO = 300
SEMENTE = 42
MODELOS = ("qwen2.5:7b", "sabiazinho-4")   # substituídos pelos nomes do cabeçalho da entrada


# Mensagens fixas da triagem a partir do M1/M2 (urgência e triagem indisponível):
# são idênticas entre si e inflariam a similaridade, então saem como as recusas.
MENSAGENS_FIXAS = ("⚠️ O que você descreveu", "Não consegui avaliar sua pergunta")


def recusa(texto: str) -> bool:
    """Bloqueio da triagem (ou resposta vazia, de execução com erro)."""
    t = texto.lstrip()
    return (not t or (t.startswith("Oi!") and "Sou especializada" in t[:80])
            or t.startswith(MENSAGENS_FIXAS))


def carregar(entrada: Path):
    global MODELOS
    ws = openpyxl.load_workbook(entrada).worksheets[1]
    rows = list(ws.iter_rows(values_only=True))
    MODELOS = tuple(str(rows[0][i]).replace("resposta", "").strip() for i in (5, 7))
    linhas = []
    for r in rows[1:]:
        if not r[0]:
            continue
        linhas.append(dict(id=r[0], cat=r[1], var=r[3], perg=r[4],
                           resp={MODELOS[0]: str(r[5] or ""), MODELOS[1]: str(r[7] or "")}))
    return linhas


_SCORER = None


def scorer():
    global _SCORER
    if _SCORER is None:
        _SCORER = BERTScorer(model_type=MODELO, num_layers=CAMADA, batch_size=8)
        # o tokenizer do BERTimbau não declara limite; o modelo aceita 512 posições
        _SCORER._tokenizer.model_max_length = 512
    return _SCORER


def truncadas(textos):
    tok = scorer()._tokenizer
    return sum(len(tok.encode(t, add_special_tokens=True)) > 512 for t in textos)


def bertscore(cands, refs):
    p, r, f = scorer().score(cands, refs)
    return p.tolist(), r.tolist(), f.tolist()


def resumo(valores):
    n = len(valores)
    media = statistics.mean(valores) if n else float("nan")
    dp = statistics.stdev(valores) if n > 1 else float("nan")
    return n, media, dp


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--entrada", type=Path, default=ENTRADA)
    ap.add_argument("--saida", type=Path, default=SAIDA)
    args = ap.parse_args()
    linhas = carregar(args.entrada)
    print("modelos:", MODELOS)
    por_id = {}
    for x in linhas:
        por_id.setdefault(x["id"], []).append(x)

    # ---- 1. pares intramodelo ------------------------------------------
    intra = []          # (modelo, id, cat, var_a, var_b, texto_a, texto_b)
    excluidos_intra = {m: 0 for m in MODELOS}
    for m in MODELOS:
        for cid, xs in por_id.items():
            for a, b in itertools.combinations(xs, 2):
                if recusa(a["resp"][m]) or recusa(b["resp"][m]):
                    excluidos_intra[m] += 1
                    continue
                intra.append((m, cid, a["cat"], a["var"], b["var"], a["resp"][m], b["resp"][m]))

    # ---- 2. pares entre modelos ----------------------------------------
    cruz = [x for x in linhas if not recusa(x["resp"][MODELOS[0]]) and not recusa(x["resp"][MODELOS[1]])]

    # ---- 3. piso: pares aleatórios de ids diferentes -------------------
    rnd = random.Random(SEMENTE)
    respondidas = [(x["id"], m, x["resp"][m]) for x in linhas for m in MODELOS if not recusa(x["resp"][m])]
    piso = []
    while len(piso) < N_PISO:
        a, b = rnd.sample(respondidas, 2)
        if a[0] != b[0]:
            piso.append((a, b))

    print(f"pares intramodelo: {len(intra)} | entre modelos: {len(cruz)} | piso: {len(piso)}")
    textos = [x[2] for x in respondidas]
    n_trunc = truncadas(textos)
    print(f"respostas acima de 512 tokens (truncadas): {n_trunc} de {len(textos)}")
    _, _, f_intra = bertscore([t[5] for t in intra], [t[6] for t in intra])
    p_c, r_c, f_cruz = bertscore([x["resp"][MODELOS[0]] for x in cruz], [x["resp"][MODELOS[1]] for x in cruz])
    _, _, f_piso = bertscore([a[2] for a, _ in piso], [b[2] for _, b in piso])

    # ---- planilha ------------------------------------------------------
    wb = openpyxl.Workbook()
    w1 = wb.active; w1.title = "Intramodelo (pares)"
    w1.append(["modelo", "id", "categoria", "variante A", "variante B", "BERTScore F1"])
    for t, f in zip(intra, f_intra):
        w1.append([t[0], t[1], t[2], t[3], t[4], round(f, 4)])

    w2 = wb.create_sheet("Entre modelos (pares)")
    w2.append(["id", "categoria", "variante", "pergunta", "P", "R", "F1"])
    for x, p, r, f in zip(cruz, p_c, r_c, f_cruz):
        w2.append([x["id"], x["cat"], x["var"], x["perg"], round(p, 4), round(r, 4), round(f, 4)])

    w3 = wb.create_sheet("Resumo")
    w3.append([f"BERTScore F1 com {MODELO}, camada {CAMADA}, sem rescala por baseline. "
               "Recusas da triagem excluídas. Desvio padrão amostral."])
    w3.append([])
    w3.append(["análise", "categoria", "n pares", "média F1", "desvio padrão"])
    cats = list(dict.fromkeys(x["cat"] for x in linhas))

    def bloco(nome, itens):     # itens: lista de (cat, f1)
        for c in cats + ["TOTAL"]:
            v = [f for cat, f in itens if c == "TOTAL" or cat == c]
            n, mu, dp = resumo(v)
            w3.append([nome, c, n, round(mu, 4) if n else None, round(dp, 4) if n > 1 else None])
            print(f"{nome:28s} {c:34s} n={n:3d} média={mu:.4f} dp={dp:.4f}" if n > 1
                  else f"{nome:28s} {c:34s} n={n:3d}")

    for m in MODELOS:
        bloco(f"intramodelo {m}", [(t[2], f) for t, f in zip(intra, f_intra) if t[0] == m])
    bloco("entre modelos", [(x["cat"], f) for x, f in zip(cruz, f_cruz)])
    n, mu, dp = resumo(f_piso)
    w3.append(["piso (ids diferentes)", "TOTAL", n, round(mu, 4), round(dp, 4)])
    print(f"{'piso (ids diferentes)':28s} {'TOTAL':34s} n={n:3d} média={mu:.4f} dp={dp:.4f}")
    w3.append([])
    for m in MODELOS:
        w3.append([f"pares intramodelo excluídos por recusa ({m})", "", excluidos_intra[m]])
    w3.append(["perguntas excluídas da comparação entre modelos", "", len(linhas) - len(cruz)])
    w3.append(["respostas truncadas em 512 tokens", "", n_trunc, f"de {len(textos)}"])
    for ws in (w1, w2, w3):
        for c in ws[1 if ws is not w3 else 3]:
            c.font = Font(bold=True)
    w2.column_dimensions["D"].width = 60
    w3.column_dimensions["A"].width = 48
    w3.column_dimensions["B"].width = 34
    wb.save(args.saida)
    print("excluídos intramodelo:", excluidos_intra, "| excluídas entre modelos:", len(linhas) - len(cruz))
    print("salvo em", args.saida)


if __name__ == "__main__":
    main()
