# -*- coding: utf-8 -*-
"""
Compara duas execuções do executar_testes.py (por exemplo, antes e depois de
uma mudança na triagem), sobre as mesmas perguntas e repetições.

Por categoria, conta: respondidas, recusas por escopo, alertas de urgência,
erros e latência média. Na execução antiga não existe a coluna "urgente":
toda pergunta bloqueada conta como recusa por escopo.

Uso:
    python comparar_execucoes.py execucoes/qwen2.5-7b.xlsx execucoes/qwen2.5-7b_m1m2.xlsx
    python comparar_execucoes.py antes.xlsx depois.xlsx --saida comparacao.xlsx
"""
import argparse
import collections
import statistics as st
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

sys.stdout.reconfigure(encoding="utf-8")


def ler(caminho: Path) -> dict:
    ws = openpyxl.load_workbook(caminho, read_only=True)["Respostas"]
    linhas = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        r = list(r) + [None] * (16 - len(r))
        if r[13]:
            desfecho = "erro"
        elif r[15] == "sim":
            desfecho = "urgente"
        elif r[8] == "sim":
            desfecho = "recusa"
        else:
            desfecho = "respondida"
        linhas[(r[0], r[3], int(r[5]))] = dict(cat=r[1], desfecho=desfecho, latencia=r[12])
    return linhas


def resumir(linhas: dict, chaves) -> dict:
    por_cat = collections.OrderedDict()
    for k in sorted(chaves, key=lambda k: (k[0][0], int(k[0][1:]), k[1], k[2])):
        x = linhas[k]
        c = por_cat.setdefault(x["cat"], dict(n=0, latencias=[], **{d: 0 for d in DESFECHOS}))
        c["n"] += 1
        c[x["desfecho"]] += 1
        if x["latencia"] is not None:
            c["latencias"].append(x["latencia"])
    return por_cat


DESFECHOS = ("respondida", "recusa", "urgente", "erro")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("antes", type=Path)
    ap.add_argument("depois", type=Path)
    ap.add_argument("--saida", type=Path)
    a = ap.parse_args()

    antes, depois = ler(a.antes), ler(a.depois)
    comuns = antes.keys() & depois.keys()
    if not comuns:
        sys.exit("As execuções não têm perguntas e repetições em comum.")
    ra, rd = resumir(antes, comuns), resumir(depois, comuns)

    cab = ["categoria", "n"] + [f"{d} ({lado})" for d in DESFECHOS for lado in ("antes", "depois")] \
          + ["latência média s (antes)", "latência média s (depois)"]
    linhas = []
    total_a = collections.Counter(); total_d = collections.Counter()
    for cat in ra:
        A, D = ra[cat], rd[cat]
        linhas.append([cat, A["n"]] + [v for d in DESFECHOS for v in (A[d], D[d])]
                      + [round(st.mean(A["latencias"]), 1) if A["latencias"] else None,
                         round(st.mean(D["latencias"]), 1) if D["latencias"] else None])
        for d in DESFECHOS + ("n",):
            total_a[d] += A[d]; total_d[d] += D[d]
    linhas.append(["TOTAL", total_a["n"]] + [v for d in DESFECHOS for v in (total_a[d], total_d[d])] + [None, None])

    # mudanças de desfecho pergunta a pergunta
    transicoes = collections.Counter((antes[k]["desfecho"], depois[k]["desfecho"]) for k in comuns)

    print(f"{len(comuns)} execuções em comum ({a.antes.name} → {a.depois.name})\n")
    larg = [34, 4] + [9] * 8
    print(" | ".join(str(c)[:w].ljust(w) for c, w in zip(["categoria", "n", "resp.ant", "resp.dep", "recusa.a",
                                                           "recusa.d", "urg.ant", "urg.dep", "erro.ant", "erro.dep"], larg)))
    for l in linhas:
        print(" | ".join(str(c if c is not None else "")[:w].ljust(w) for c, w in zip(l, larg)))
    print("\nTransições de desfecho (antes → depois):")
    for (x, y), n in sorted(transicoes.items(), key=lambda t: -t[1]):
        print(f"  {x:>10} → {y:<10} {n}")

    if a.saida:
        wb = openpyxl.Workbook()
        w = wb.active; w.title = "Por categoria"
        w.append(cab)
        for l in linhas:
            w.append(l)
        t = wb.create_sheet("Transições")
        t.append(["antes", "depois", "n"])
        for (x, y), n in sorted(transicoes.items(), key=lambda t: -t[1]):
            t.append([x, y, n])
        d = wb.create_sheet("Por pergunta")
        d.append(["id", "variante", "repeticao", "categoria", "antes", "depois"])
        for k in sorted(comuns, key=lambda k: (k[0][0], int(k[0][1:]), k[1], k[2])):
            d.append([k[0], k[1], k[2], antes[k]["cat"], antes[k]["desfecho"], depois[k]["desfecho"]])
        for ws in (w, t, d):
            for c in ws[1]:
                c.font = Font(bold=True)
        wb.save(a.saida)
        print("\nsalvo em", a.saida)


if __name__ == "__main__":
    main()
