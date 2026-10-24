# -*- coding: utf-8 -*-
"""
BERTScore da reavaliação (Subseção 5.4): linha de base × depois das melhorias.

Lê duas ou mais execuções de `executar_testes.py executar` (aba "Respostas", um
modelo, N repetições) e calcula, para cada uma, com o mesmo BERTimbau/camada do
`bertscore_testes_adversariais.py`:

  1. Intramodelo: dentro de cada repetição, as respostas às variantes de um
     cenário duas a duas (como na avaliação original), juntando as repetições.
  2. Entre repetições: a mesma pergunta em repetições diferentes (r1×r2, r1×r3,
     r2×r3) — estabilidade do sistema diante da mesma entrada.
  3. Piso: 300 pares aleatórios de respostas a cenários diferentes.
  4. Diferença F × A–E no intramodelo (critério da M4), por permutação.

Entre execuções, compara o intramodelo total por permutação e também restrito aos
pares que existem nas duas (mesmo id, variantes e repetição, ambos respondidos),
para separar o efeito das respostas que passaram a existir (menos recusas).

Ficam de fora as recusas e todas as mensagens fixas (urgência, triagem
indisponível, sem evidência, falha de idioma), que inflariam a similaridade.

Uso (venv com torch + bert-score + openpyxl):
    python bertscore_reavaliacao.py execucoes/qwen2.5-7b.xlsx execucoes/qwen2.5-7b_m5-m7.xlsx \
        --saida bertscore_reavaliacao.xlsx
"""
import argparse, itertools, random, statistics
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Font

import bertscore_testes_adversariais as bs

N_PISO = 300
N_PERM = 20000
SEMENTE = 42

# Mensagens fixas do agente de recuperação (M5, M7), além das da triagem.
FIXAS_RETRIEVAL = ("Não encontrei no PubMed", "Não consegui gerar a resposta em português")


def recusa(texto: str) -> bool:
    t = (texto or "").lstrip()
    return bs.recusa(t) or t.startswith(FIXAS_RETRIEVAL)


def carregar(caminho: Path):
    ws = openpyxl.load_workbook(caminho, read_only=True)["Respostas"]
    rows = list(ws.iter_rows(values_only=True))
    h = rows[0]
    linhas = []
    for r in rows[1:]:
        d = dict(zip(h, r))
        if not d.get("id") or d.get("erro"):
            continue
        linhas.append(dict(id=d["id"], cat=d["categoria"], var=d["variante"],
                           rep=int(d["repeticao"]), resp=str(d["resposta"] or "")))
    return linhas


def permutacao(a, b, rnd):
    """p bilateral da diferença de médias entre a e b."""
    gen = np.random.default_rng(rnd.randrange(2**32))
    juntos, n = np.array(a + b), len(a)
    obs = abs(juntos[:n].mean() - juntos[n:].mean())
    extremos = 0
    for _ in range(N_PERM):
        gen.shuffle(juntos)
        if abs(juntos[:n].mean() - juntos[n:].mean()) >= obs - 1e-12:
            extremos += 1
    return (extremos + 1) / (N_PERM + 1)


def analisar(nome, linhas, rnd):
    resp = [x for x in linhas if not recusa(x["resp"])]
    por_id_rep, por_pergunta = {}, {}
    for x in resp:
        por_id_rep.setdefault((x["id"], x["rep"]), []).append(x)
        por_pergunta.setdefault((x["id"], x["var"]), []).append(x)

    intra = []      # (cat, id, var_a, var_b, rep, texto_a, texto_b)
    for (cid, rep), xs in por_id_rep.items():
        for a, b in itertools.combinations(sorted(xs, key=lambda x: x["var"]), 2):
            intra.append((a["cat"], cid, a["var"], b["var"], rep, a["resp"], b["resp"]))
    entre_rep = []  # (cat, id, var, rep_a, rep_b, texto_a, texto_b)
    for (cid, var), xs in por_pergunta.items():
        for a, b in itertools.combinations(sorted(xs, key=lambda x: x["rep"]), 2):
            entre_rep.append((a["cat"], cid, var, a["rep"], b["rep"], a["resp"], b["resp"]))
    piso = []
    while len(piso) < N_PISO:
        a, b = rnd.sample(resp, 2)
        if a["id"] != b["id"]:
            piso.append((a["resp"], b["resp"]))

    print(f"\n{nome}: {len(resp)} respostas | intramodelo {len(intra)} | "
          f"entre repetições {len(entre_rep)} | piso {len(piso)}")
    n_trunc = bs.truncadas([x["resp"] for x in resp])
    _, _, f_intra = bs.bertscore([t[5] for t in intra], [t[6] for t in intra])
    _, _, f_rep = bs.bertscore([t[5] for t in entre_rep], [t[6] for t in entre_rep])
    _, _, f_piso = bs.bertscore([a for a, _ in piso], [b for _, b in piso])
    return dict(nome=nome, n_resp=len(resp), n_trunc=n_trunc, intra=list(zip(intra, f_intra)),
                entre_rep=list(zip(entre_rep, f_rep)), piso=f_piso)


def fmt(v):
    n, mu, dp = bs.resumo(v)
    return f"n={n:3d} média={mu:.4f} dp={dp:.4f}" if n > 1 else f"n={n:3d}"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("execucoes", type=Path, nargs="+")
    ap.add_argument("--saida", type=Path, default=Path(__file__).with_name("bertscore_reavaliacao.xlsx"))
    args = ap.parse_args()
    rnd = random.Random(SEMENTE)

    res = [analisar(p.stem, carregar(p), rnd) for p in args.execucoes]
    cats = sorted({t[0] for r in res for t, _ in r["intra"]})

    wb = openpyxl.Workbook()
    w = wb.active; w.title = "Resumo"
    w.append([f"BERTScore F1 com {bs.MODELO}, camada {bs.CAMADA}, sem rescala. Recusas e "
              "mensagens fixas excluídas. Desvio padrão amostral; p por permutação "
              f"({N_PERM} reamostragens, bilateral)."])
    w.append([])
    w.append(["execução", "análise", "categoria", "n pares", "média F1", "desvio padrão"])
    for r in res:
        print(f"\n== {r['nome']}  (respostas truncadas em 512 tokens: {r['n_trunc']} de {r['n_resp']})")
        for analise in ("intra", "entre_rep"):
            for c in cats + ["TOTAL"]:
                v = [f for t, f in r[analise] if c == "TOTAL" or t[0] == c]
                n, mu, dp = bs.resumo(v)
                w.append([r["nome"], analise, c, n, round(mu, 4) if n else None,
                          round(dp, 4) if n > 1 else None])
                print(f"  {analise:9s} {c:34s} {fmt(v)}")
        n, mu, dp = bs.resumo(r["piso"])
        w.append([r["nome"], "piso", "TOTAL", n, round(mu, 4), round(dp, 4)])
        print(f"  {'piso':9s} {'TOTAL':34s} {fmt(r['piso'])}")
        f_ = [f for t, f in r["intra"] if t[0].startswith("F")]
        ae = [f for t, f in r["intra"] if not t[0].startswith("F")]
        p = permutacao(f_, ae, rnd)
        dif = statistics.mean(f_) - statistics.mean(ae)
        w.append([r["nome"], "intra F − A–E", "", "", round(dif, 4), f"p = {p:.4f}"])
        print(f"  intra F − A–E: {dif:+.4f} (p = {p:.4f})")
        w.append([r["nome"], "respostas truncadas em 512 tokens", "", r["n_trunc"], f"de {r['n_resp']}"])

    # Comparação entre a primeira execução (antes) e cada uma das demais.
    w.append([])
    w.append(["comparação", "análise", "", "n", "diferença (depois − antes)", "p"])
    antes = res[0]
    for depois in res[1:]:
        for analise in ("intra", "entre_rep"):
            a = [f for _, f in antes[analise]]
            b = [f for _, f in depois[analise]]
            p = permutacao(a, b, rnd)
            dif = statistics.mean(b) - statistics.mean(a)
            w.append([f"{antes['nome']} → {depois['nome']}", analise, "", f"{len(a)} × {len(b)}",
                      round(dif, 4), round(p, 4)])
            print(f"\n{antes['nome']} → {depois['nome']} {analise}: {dif:+.4f} (p = {p:.4f})")
        chave = lambda t: (t[1], t[2], t[3], t[4])
        fa = {chave(t): f for t, f in antes["intra"]}
        fb = {chave(t): f for t, f in depois["intra"]}
        comuns = sorted(fa.keys() & fb.keys())
        a = [fa[k] for k in comuns]
        b = [fb[k] for k in comuns]
        p = permutacao(a, b, rnd)
        dif = statistics.mean(b) - statistics.mean(a)
        w.append([f"{antes['nome']} → {depois['nome']}", "intra, só pares em comum", "", len(comuns),
                  round(dif, 4), round(p, 4)])
        print(f"{antes['nome']} → {depois['nome']} intra em comum (n={len(comuns)}): "
              f"{statistics.mean(a):.4f} → {statistics.mean(b):.4f} ({dif:+.4f}, p = {p:.4f})")

    for r in res:
        ws = wb.create_sheet(f"Pares {r['nome']}"[:31])
        ws.append(["análise", "categoria", "id", "variante/rep A", "variante/rep B", "repetição/variante", "F1"])
        for t, f in r["intra"]:
            ws.append(["intra", t[0], t[1], t[2], t[3], t[4], round(f, 4)])
        for t, f in r["entre_rep"]:
            ws.append(["entre_rep", t[0], t[1], t[3], t[4], t[2], round(f, 4)])
        for c in ws[1]:
            c.font = Font(bold=True)
    for c in w[3]:
        c.font = Font(bold=True)
    w.column_dimensions["A"].width = 34
    w.column_dimensions["B"].width = 26
    w.column_dimensions["C"].width = 34
    wb.save(args.saida)
    print("\nsalvo em", args.saida)


if __name__ == "__main__":
    main()
